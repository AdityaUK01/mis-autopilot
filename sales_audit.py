#!/usr/bin/env python3
"""Check a sales workbook and write a small Markdown audit report."""

from __future__ import annotations

import argparse
import math
import re
import sys
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ALIASES = {
    "order_id": {"order id", "order number", "invoice number", "invoice no", "transaction id", "sales id"},
    "date": {"date", "order date", "sale date", "invoice date", "transaction date"},
    "product": {"product", "product name", "item", "item name", "sku", "description"},
    "quantity": {"quantity", "qty", "units", "unit sold", "units sold"},
    "unit_price": {"unit price", "price", "selling price", "sale price", "rate"},
    "total": {"total", "sales", "sales amount", "total amount", "line total", "revenue", "amount"},
}


def normalize(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").strip().lower()).strip()


def number(value: Any) -> float | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        result = float(value)
        return result if math.isfinite(result) else None
    raw = str(value).strip().replace(",", "")
    raw = re.sub(r"^[^\d+\-.]+", "", raw)
    raw = re.sub(r"[^\d.\-]+$", "", raw)
    try:
        result = float(raw)
        return result if math.isfinite(result) else None
    except ValueError:
        return None


def nonblank(value: Any) -> bool:
    return value is not None and (not isinstance(value, str) or bool(value.strip()))


def find_columns(headers: list[Any]) -> dict[str, int]:
    normalized = [normalize(item) for item in headers]
    result: dict[str, int] = {}
    for field, names in ALIASES.items():
        for index, header in enumerate(normalized):
            if header in names:
                result[field] = index
                break
    return result


def audit_sheet(sheet: Any, selected_sheet: str) -> tuple[str, int, list[tuple[int, str, str]]]:
    rows = sheet.iter_rows(values_only=True)
    header_row = None
    columns: dict[str, int] = {}
    header_index = 0
    # Look at the first 15 rows to tolerate a title or blank rows above the table.
    for index, values in enumerate(rows):
        candidate = find_columns(list(values))
        if len(candidate) >= 2:
            header_row, columns, header_index = values, candidate, index
            break
        if index >= 14:
            break

    if header_row is None:
        return (f"## Sheet: {selected_sheet}\n\nCould not find a header row with recognizable sales columns.\n", 0, [])

    headers = [str(v).strip() if v is not None else f"Column {i + 1}" for i, v in enumerate(header_row)]
    missing_columns = [field for field in ("order_id", "date", "product", "quantity", "unit_price", "total") if field not in columns]
    issues: list[tuple[int, str, str]] = []
    row_count = 0
    totals: list[float] = []
    qty_idx, price_idx, total_idx = (columns.get(key) for key in ("quantity", "unit_price", "total"))

    for row_number, values in enumerate(rows, start=header_index + 2):
        if not any(nonblank(value) for value in values):
            continue
        row_count += 1
        cells = list(values)

        for field, col in columns.items():
            value = cells[col] if col < len(cells) else None
            if not nonblank(value):
                issues.append((row_number, f"Missing {field.replace('_', ' ')}", headers[col]))
                continue
            if field == "date" and not isinstance(value, (date, datetime)):
                try:
                    datetime.fromisoformat(str(value).strip().replace("Z", "+00:00"))
                except ValueError:
                    issues.append((row_number, "Invalid date", str(value)))
            elif field in {"quantity", "unit_price", "total"}:
                parsed = number(value)
                if parsed is None:
                    issues.append((row_number, f"Invalid {field.replace('_', ' ')}", str(value)))
                elif field == "quantity" and parsed <= 0:
                    issues.append((row_number, "Quantity must be greater than 0", str(value)))
                elif field in {"unit_price", "total"} and parsed < 0:
                    issues.append((row_number, f"{field.replace('_', ' ').title()} cannot be negative", str(value)))

        if total_idx is not None:
            parsed_total = number(cells[total_idx] if total_idx < len(cells) else None)
            if parsed_total is not None:
                totals.append(parsed_total)
        if qty_idx is not None and price_idx is not None and total_idx is not None:
            qty = number(cells[qty_idx] if qty_idx < len(cells) else None)
            price = number(cells[price_idx] if price_idx < len(cells) else None)
            total = number(cells[total_idx] if total_idx < len(cells) else None)
            if qty is not None and price is not None and total is not None and not math.isclose(qty * price, total, rel_tol=0.01, abs_tol=0.02):
                issues.append((row_number, "Total does not match quantity × unit price", f"{total:g} vs {qty * price:g}"))

    counts = Counter(reason for _, reason, _ in issues)
    lines = [f"## Sheet: {selected_sheet}", "", f"- Data rows checked: **{row_count}**", f"- Issues found: **{len(issues)}**"]
    if totals:
        lines.append(f"- Sum of recognized total column: **{sum(totals):,.2f}**")
    if missing_columns:
        lines += ["", "Recognized columns: " + ", ".join(f"`{field}` → `{headers[col]}`" for field, col in columns.items()),
                  "Missing columns (checks requiring these were skipped): " + ", ".join(f"`{name}`" for name in missing_columns)]
    if counts:
        lines += ["", "### Issue counts", ""]
        lines.extend(f"- {reason}: {count}" for reason, count in sorted(counts.items()))
        lines += ["", "### Details", "", "| Row | Issue | Value |", "|---:|---|---|"]
        for row, reason, value in issues:
            safe_value = str(value).replace("|", "\\|").replace("\n", " ")
            lines.append(f"| {row} | {reason} | {safe_value} |")
    else:
        lines += ["", "No issues were found in the recognized columns."]
    return "\n".join(lines) + "\n", row_count, issues


def main() -> int:
    parser = argparse.ArgumentParser(description="Check a sales Excel workbook and write a Markdown summary report.")
    parser.add_argument("workbook", type=Path, help="Path to an .xlsx or .xlsm sales workbook")
    parser.add_argument("--sheet", help="Sheet name to check (default: first sheet)")
    parser.add_argument("--output", type=Path, help="Report path (default: <workbook>_report.md)")
    args = parser.parse_args()
    if not args.workbook.is_file():
        parser.error(f"Workbook not found: {args.workbook}")
    if args.workbook.suffix.lower() not in {".xlsx", ".xlsm"}:
        parser.error("Only .xlsx and .xlsm workbooks are supported.")
    output = args.output or args.workbook.with_name(f"{args.workbook.stem}_report.md")
    try:
        book = load_workbook(args.workbook, read_only=True, data_only=True)
    except Exception as exc:
        print(f"Could not open workbook: {exc}", file=sys.stderr)
        return 2
    if args.sheet and args.sheet not in book.sheetnames:
        print(f"Sheet not found: {args.sheet}. Available sheets: {', '.join(book.sheetnames)}", file=sys.stderr)
        return 2
    sheet_names = [args.sheet] if args.sheet else [book.sheetnames[0]]
    reports = [audit_sheet(book[name], name)[0] for name in sheet_names]
    report = f"# Sales workbook check\n\nSource: `{args.workbook.name}`\n\n" + "\n".join(reports)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(report, encoding="utf-8")
    print(f"Report written to {output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
