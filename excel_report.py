"""Build a formatted, auditable Excel MIS workbook for the current analysis."""

from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from typing import Any

import pandas as pd
from mis_core import validate_sales_result_schema
from report_templates import validate_template_result_schema
from spreadsheet_core import validate_profile_schema
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties


NAVY = "17324D"
TEAL = "167B72"
WHITE = "FFFFFF"
PALE_BLUE = "EAF2F8"
PALE_RED = "FCE4E4"
PALE_AMBER = "FFF2CC"
PALE_GREEN = "E2F0D9"
PALE_GRAY = "F2F4F7"
TEXT = "243447"
INR_FORMAT = '₹#,##0.00;[Red]-₹#,##0.00'
INTEGER_FORMAT = "#,##0"
THIN_RULE = Side(style="thin", color="D8E1EA")
BODY_FONT = Font(name="Aptos", size=10, color=TEXT)
BODY_ALIGNMENT = Alignment(vertical="top")
BODY_WRAP_ALIGNMENT = Alignment(vertical="top", wrap_text=True)
BODY_BORDER = Border(bottom=Side(style="hair", color="E6EBF0"))
STATUS_FILLS = {
    "Valid": PatternFill("solid", fgColor=PALE_GREEN),
    "Needs review": PatternFill("solid", fgColor=PALE_RED),
    "Blank row": PatternFill("solid", fgColor=PALE_GRAY),
}


def _blank(value: Any) -> bool:
    if value is None:
        return True
    try:
        if pd.isna(value):
            return True
    except (TypeError, ValueError):
        pass
    return isinstance(value, str) and not value.strip()


def _excel_value(value: Any) -> Any:
    """Convert pandas/numpy scalars without coercing original cell contents."""
    if _blank(value):
        return None
    if isinstance(value, pd.Timestamp):
        if value.tzinfo is not None:
            return value.isoformat()
        return value.to_pydatetime()
    if hasattr(value, "item") and not isinstance(value, (str, bytes, date, datetime)):
        try:
            value = value.item()
        except (ValueError, AttributeError):
            pass
    if isinstance(value, datetime) and value.tzinfo is not None:
        return value.isoformat()
    if isinstance(value, (str, int, float, bool, date, datetime)):
        return value
    return str(value)


def _write_title(ws: Any, title: str, subtitle: str, final_column: int) -> None:
    final_letter = get_column_letter(max(final_column, 1))
    ws.merge_cells(f"A1:{final_letter}1")
    title_cell = ws["A1"]
    title_cell.value = title
    title_cell.fill = PatternFill("solid", fgColor=NAVY)
    title_cell.font = Font(name="Aptos Display", size=16, bold=True, color=WHITE)
    title_cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 30
    ws.merge_cells(f"A2:{final_letter}2")
    note_cell = ws["A2"]
    note_cell.value = subtitle
    note_cell.font = Font(name="Aptos", size=10, italic=True, color=TEXT)
    note_cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 34
    ws.sheet_view.showGridLines = False


def _write_headers(ws: Any, headers: list[str], row: int = 4) -> None:
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.fill = PatternFill("solid", fgColor=TEAL)
        cell.font = Font(name="Aptos", size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = Border(bottom=Side(style="medium", color=NAVY))
    ws.row_dimensions[row].height = 32
    ws.freeze_panes = f"A{row + 1}"


def _fit_widths(ws: Any, min_row: int, max_row: int, min_col: int, max_col: int, cap: int = 48) -> None:
    for col in range(min_col, max_col + 1):
        lengths = [len(str(ws.cell(row=row, column=col).value or "")) for row in range(min_row, max_row + 1)]
        ws.column_dimensions[get_column_letter(col)].width = min(max(max(lengths, default=8) + 2, 12), cap)


def _write_summary_sheet(ws: Any, title: str, subtitle: str, group_label: str, summary: list[dict[str, Any]]) -> int:
    headers = [group_label, "Raw Records", "Validated Records", "Raw Sales (INR)", "Validated Sales (INR)"]
    _write_title(ws, title, subtitle, len(headers))
    _write_headers(ws, headers)
    first_data_row = 5
    for row_num, record in enumerate(summary, start=first_data_row):
        values = [record["Label"], record["Raw Records"], record["Validated Records"], record["Raw Sales"], record["Validated Sales"]]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row_num, column=col, value=_excel_value(value))
            cell.font = Font(name="Aptos", size=10, color=TEXT)
            cell.alignment = Alignment(vertical="top", wrap_text=(col == 1))
            cell.border = Border(bottom=Side(style="hair", color="E6EBF0"))
            if col in (4, 5):
                cell.number_format = INR_FORMAT
            elif col in (2, 3):
                cell.number_format = INTEGER_FORMAT
    last_data_row = first_data_row + len(summary) - 1
    total_row = last_data_row + 1
    ws.cell(total_row, 1, "TOTAL")
    for col in range(2, 6):
        letter = get_column_letter(col)
        if summary:
            ws.cell(total_row, col, f"=SUM({letter}{first_data_row}:{letter}{last_data_row})")
        else:
            ws.cell(total_row, col, 0)
    for cell in ws[total_row]:
        if cell.column <= len(headers):
            cell.fill = PatternFill("solid", fgColor=PALE_BLUE)
            cell.font = Font(name="Aptos", bold=True, color=NAVY)
            cell.border = Border(top=Side(style="medium", color=TEAL))
            if cell.column in (4, 5):
                cell.number_format = INR_FORMAT
            elif cell.column in (2, 3):
                cell.number_format = INTEGER_FORMAT
    ws.auto_filter.ref = f"A4:{get_column_letter(len(headers))}{max(last_data_row, 4)}"
    _fit_widths(ws, 4, max(total_row, 4), 1, len(headers))
    ws.column_dimensions["A"].width = min(max(ws.column_dimensions["A"].width, 24), 48)
    return total_row


def _add_comparison_chart(ws: Any, source_ws: Any, title: str, group_label: str, anchor: str, max_rows: int = 10) -> None:
    data_count = max(0, source_ws.max_row - 5)  # exclude title, header and total rows
    count = min(data_count, max_rows)
    if count <= 0:
        return
    first_data_row = 5
    last_data_row = first_data_row + count - 1
    chart = BarChart()
    chart.type = "bar"
    chart.style = 10
    chart.grouping = "clustered"
    chart.overlap = 0
    chart.title = title
    chart.x_axis.title = "Sales (INR)"
    chart.y_axis.title = group_label
    chart.x_axis.numFmt = INR_FORMAT
    chart.height = 8.0
    chart.width = 15.5
    chart.add_data(Reference(source_ws, min_col=4, max_col=5, min_row=4, max_row=last_data_row), titles_from_data=True)
    chart.set_categories(Reference(source_ws, min_col=1, min_row=first_data_row, max_row=last_data_row))
    chart.legend.position = "b"
    ws.add_chart(chart, anchor)


def _add_trend_chart(ws: Any, first_data_row: int, last_data_row: int) -> None:
    if last_data_row < first_data_row:
        return
    chart = LineChart()
    chart.style = 13
    chart.title = "Sales trend (valid dates)"
    chart.x_axis.title = "Date / month"
    chart.y_axis.title = "Sales (INR)"
    chart.y_axis.numFmt = INR_FORMAT
    chart.height = 8.0
    chart.width = 24.0
    chart.add_data(Reference(ws, min_col=2, max_col=3, min_row=first_data_row - 1, max_row=last_data_row), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=first_data_row, max_row=last_data_row))
    chart.legend.position = "b"
    ws.add_chart(chart, "A31")


def build_excel_report(result: dict[str, Any]) -> bytes:
    """Return a formatted Excel MIS workbook as bytes for Streamlit download."""
    validate_sales_result_schema(result)
    workbook = Workbook()
    dashboard = workbook.active
    dashboard.title = "MIS Dashboard"
    sales_ws = workbook.create_sheet("Sales Data")
    issues_ws = workbook.create_sheet("Data Quality Issues")
    product_ws = workbook.create_sheet("Product Summary")
    region_ws = workbook.create_sheet("Region Summary") if result["region_available"] else None
    workbook.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True)
    workbook.properties.title = "MIS Autopilot Sales Report"
    workbook.properties.subject = "Sales MIS and data-quality review"
    workbook.properties.creator = "MIS Autopilot"

    source_data: pd.DataFrame = result["original_data"]
    source_headers = [str(column) for column in source_data.columns]
    row_note = (
        "Original source values are preserved below, including blank rows. Status and issue columns are added for review. "
        "Validated summaries exclude exact full-row duplicates only. Repeated transaction candidates that differ only by Record ID stay included until reviewed."
    )
    _write_title(sales_ws, "Sales Data | Original Source", f"Source: {result['source_name']} · {row_note}", len(source_headers) + 4)
    duplicate_header = "Exact Duplicate"
    repeated_header = "Repeated Transaction Candidate"
    status_header, flags_header = "Quality Status", "Issue Indicators"
    headers = source_headers + [duplicate_header, repeated_header, status_header, flags_header]
    _write_headers(sales_ws, headers)
    source_col_letters = {str(header): get_column_letter(index + 1) for index, header in enumerate(source_data.columns)}
    mapping = result["mapping"]
    issue_rows: dict[int, list[dict[str, Any]]] = {}
    for record in result["issues"].to_dict("records"):
        issue_rows.setdefault(int(record["Row"]), []).append(record)
    duplicate_rows = {
        int(record["Row"]) for record in result["issues"].to_dict("records") if record["Issue"] == "Duplicate record"
    }
    repeated_transaction_rows = {
        int(record["Row"]) for record in result["issues"].to_dict("records") if record["Issue"] == "Repeated transaction candidate"
    }
    data_first_row = 5
    duplicate_col = len(source_headers) + 1
    repeated_col = len(source_headers) + 2
    status_col = len(source_headers) + 3
    flags_col = len(source_headers) + 4
    for index, (_, source_row) in enumerate(source_data.iterrows()):
        excel_row = data_first_row + index
        source_row_number = index + 2
        row_values = [_excel_value(source_row[column]) for column in source_data.columns]
        row_issues = issue_rows.get(source_row_number, [])
        source_nonblank = any(not _blank(value) for value in source_row.tolist())
        status = "Needs review" if row_issues else ("Valid" if source_nonblank else "Blank row")
        flags = "; ".join(dict.fromkeys(issue["Issue"] for issue in row_issues)) if row_issues else ("No issues" if source_nonblank else "")
        values = row_values + [
            "Yes" if source_row_number in duplicate_rows else "No",
            "Yes" if source_row_number in repeated_transaction_rows else "No",
            status,
            flags,
        ]
        for col, value in enumerate(values, start=1):
            cell = sales_ws.cell(excel_row, col, value)
            cell.font = BODY_FONT
            cell.alignment = BODY_WRAP_ALIGNMENT if col > len(source_headers) else BODY_ALIGNMENT
            cell.border = BODY_BORDER
        sales_ws.cell(excel_row, status_col).fill = STATUS_FILLS[status]

    last_sales_row = max(data_first_row + len(source_data) - 1, data_first_row)
    sales_last_col = get_column_letter(len(headers))
    sales_ws.auto_filter.ref = f"A4:{sales_last_col}{last_sales_row}"
    _fit_widths(sales_ws, 4, max(last_sales_row, 4), 1, len(headers), cap=44)
    for col in (duplicate_col, repeated_col, status_col):
        sales_ws.column_dimensions[get_column_letter(col)].width = 18
    sales_ws.column_dimensions[get_column_letter(flags_col)].width = 42
    sales_ws.conditional_formatting.add(
        f"A{data_first_row}:{get_column_letter(len(source_headers))}{last_sales_row}",
        FormulaRule(formula=[f'${get_column_letter(duplicate_col)}{data_first_row}="Yes"'], fill=PatternFill("solid", fgColor=PALE_AMBER)),
    )
    sales_ws.conditional_formatting.add(
        f"A{data_first_row}:{get_column_letter(len(source_headers))}{last_sales_row}",
        FormulaRule(formula=[f'${get_column_letter(repeated_col)}{data_first_row}="Yes"'], fill=PatternFill("solid", fgColor=PALE_AMBER)),
    )
    sales_ws.conditional_formatting.add(
        f"{get_column_letter(status_col)}{data_first_row}:{get_column_letter(status_col)}{last_sales_row}",
        FormulaRule(formula=[f'${get_column_letter(status_col)}{data_first_row}="Valid"'], fill=PatternFill("solid", fgColor=PALE_GREEN)),
    )
    sales_ws.conditional_formatting.add(
        f"{get_column_letter(status_col)}{data_first_row}:{get_column_letter(status_col)}{last_sales_row}",
        FormulaRule(formula=[f'${get_column_letter(status_col)}{data_first_row}="Needs review"'], fill=PatternFill("solid", fgColor=PALE_RED)),
    )
    issue_indicator_letter = get_column_letter(flags_col)
    reasons_by_field = {
        "date": ["Missing date", "Invalid date"],
        "product": ["Missing product"],
        "quantity": ["Missing quantity", "Invalid quantity", "Negative quantity"],
        "unit_price": ["Missing unit price", "Invalid unit price", "Negative unit price"],
        "total_sales": ["Missing total sales", "Invalid total sales", "Negative total sales", "Total sales does not match quantity × unit price"],
        "region": ["Missing region"],
        "order_id": ["Missing order id"],
    }
    for field, reasons in reasons_by_field.items():
        source_column = mapping.get(field)
        if source_column not in source_col_letters:
            continue
        fill = PALE_AMBER if field in {"region", "order_id"} else PALE_RED
        for reason in reasons:
            formula = f'ISNUMBER(SEARCH("{reason}",${issue_indicator_letter}{data_first_row}))'
            sales_ws.conditional_formatting.add(
                f"{source_col_letters[source_column]}{data_first_row}:{source_col_letters[source_column]}{last_sales_row}",
                FormulaRule(formula=[formula], fill=PatternFill("solid", fgColor=fill)),
            )
    for field in ("unit_price", "total_sales"):
        source_column = mapping.get(field)
        if source_column in source_col_letters:
            col_letter = source_col_letters[source_column]
            for cell in sales_ws[f"{col_letter}{data_first_row}:{col_letter}{last_sales_row}"][0]:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = INR_FORMAT
    if mapping.get("date") in source_col_letters:
        date_letter = source_col_letters[mapping["date"]]
        for cell in sales_ws[f"{date_letter}{data_first_row}:{date_letter}{last_sales_row}"][0]:
            if isinstance(cell.value, (date, datetime)):
                cell.number_format = "dd-mmm-yyyy"

    issue_headers = ["Source Row", "Column Name", "Issue Type", "Original Value", "Explanation", "Recommended Action", "Severity"]
    _write_title(issues_ws, "Data Quality Issues", "Each finding points to the original source row; no flagged source record is removed from Sales Data.", len(issue_headers))
    _write_headers(issues_ws, issue_headers)
    for row_num, record in enumerate(result["issues"].to_dict("records"), start=5):
        values = [record["Row"], record["Column"], record["Issue"], record["Value"], record["Explanation"], record["Recommended Action"], record["Severity"]]
        for col, value in enumerate(values, start=1):
            cell = issues_ws.cell(row_num, col, _excel_value(value))
            cell.font = BODY_FONT
            cell.alignment = BODY_WRAP_ALIGNMENT if col in {4, 5, 6} else BODY_ALIGNMENT
            cell.border = BODY_BORDER
        issues_ws.row_dimensions[row_num].height = 48
    issue_last_row = max(4 + len(result["issues"]), 4)
    issues_ws.auto_filter.ref = f"A4:G{issue_last_row}"
    _fit_widths(issues_ws, 4, max(issue_last_row, 4), 1, 7, cap=48)
    for col, width in {1: 14, 2: 20, 3: 48, 4: 32, 5: 58, 6: 58, 7: 15}.items():
        issues_ws.column_dimensions[get_column_letter(col)].width = width
    if issue_last_row >= 5:
        issues_ws.conditional_formatting.add(
            f"A5:G{issue_last_row}",
            FormulaRule(formula=['$G5="Critical"'], fill=PatternFill("solid", fgColor=PALE_RED)),
        )
        issues_ws.conditional_formatting.add(
            f"A5:G{issue_last_row}",
            FormulaRule(formula=['$G5="Warning"'], fill=PatternFill("solid", fgColor=PALE_AMBER)),
        )
    else:
        issues_ws["A5"] = "No data-quality issues found."
        issues_ws["A5"].font = Font(name="Aptos", color=TEAL, bold=True)

    duplicate_note = (
        f"Raw includes all {result['record_count']:,} nonblank source records, including {result['duplicate_count']:,} exact full-row duplicate(s) and "
        f"{result['repeated_transaction_count']:,} repeated transaction candidate row(s) matching all other columns except Record ID. "
        "Validated excludes exact full-row duplicates only; repeated transaction candidates and other nonduplicate records remain included when parseable. "
        "Unparseable or missing sales amounts are omitted from numeric totals and remain listed as issues."
    )
    product_total_row = _write_summary_sheet(product_ws, "Product Summary", duplicate_note, "Product", result["product_summary"])
    region_total_row = None
    if region_ws is not None:
        region_total_row = _write_summary_sheet(region_ws, "Region Summary", duplicate_note, "Region", result["region_summary"])

    _write_title(dashboard, "MIS Autopilot | Sales Dashboard", f"Source: {result['source_name']} · Generated from the currently selected dataset.", 20)
    dashboard.merge_cells("A3:H3")
    dashboard["A3"] = (
        "Raw totals include every nonblank row. Validated totals exclude exact full-row duplicates only; repeated transaction candidates remain included pending review. "
        "Trend charts use valid dates only."
    )
    dashboard["A3"].font = Font(name="Aptos", size=10, color=TEXT, italic=True)
    dashboard["A3"].alignment = Alignment(wrap_text=True, vertical="top")
    dashboard.row_dimensions[3].height = 32
    for col, value in enumerate(["Metric", "Value", "Basis / note"], start=1):
        cell = dashboard.cell(5, col, value)
        cell.fill = PatternFill("solid", fgColor=TEAL)
        cell.font = Font(name="Aptos", bold=True, color=WHITE)
        cell.alignment = Alignment(vertical="center")
    metrics = [
        ("Raw total sales", f"='Product Summary'!D{product_total_row}", "All nonblank records; duplicates included."),
        ("Validated total sales", f"='Product Summary'!E{product_total_row}", "Exact duplicates excluded; other records retained."),
        ("Unique orders", result["unique_order_count"] if result["unique_order_count"] is not None else "Not available", "Uses distinct nonblank order IDs." if result["unique_order_count"] is not None else "No order ID column was recognized."),
        ("Average order value", result["average_order_value"], "Validated sales / unique orders; per-record average if no order ID."),
        ("Total records", result["record_count"], "Nonblank source rows, including exact duplicates."),
        ("Exact duplicate rows", result["duplicate_count"], "Full-row matches; excluded from validated summaries."),
        ("Repeated transaction candidates", result["repeated_transaction_count"], "All other source columns match; Record ID differs. Retained pending review."),
        ("Data-quality issue count", len(result["issues"]), "Counts individual findings; one row may have multiple findings."),
    ]
    metric_rows: dict[str, int] = {}
    for row_num, (label, value, note) in enumerate(metrics, start=6):
        metric_rows[label] = row_num
        dashboard.cell(row_num, 1, label)
        dashboard.cell(row_num, 2, value)
        dashboard.cell(row_num, 3, note)
        for cell in dashboard[row_num][:3]:
            cell.font = Font(name="Aptos", size=10, color=TEXT)
            cell.alignment = Alignment(vertical="top", wrap_text=(cell.column == 3))
            cell.border = Border(bottom=Side(style="thin", color="E6EBF0"))
        if label in {"Raw total sales", "Validated total sales", "Average order value"}:
            dashboard.cell(row_num, 2).number_format = INR_FORMAT
        elif isinstance(value, int):
            dashboard.cell(row_num, 2).number_format = INTEGER_FORMAT
    dashboard.column_dimensions["A"].width = 30
    dashboard.column_dimensions["B"].width = 22
    dashboard.column_dimensions["C"].width = 68
    for col in range(4, 21):
        dashboard.column_dimensions[get_column_letter(col)].width = 12
    dashboard.freeze_panes = "A6"
    if result["product_summary"]:
        _add_comparison_chart(dashboard, product_ws, "Sales by product | raw vs validated", "Product", "A15")
    if region_ws is not None and result["region_summary"]:
        _add_comparison_chart(dashboard, region_ws, "Sales by region | raw vs validated", "Region", "K15", max_rows=15)
    if result["sales_trend"]:
        trend_header_row = 44
        trend_data_row = trend_header_row + 1
        for col, header in enumerate(["Date / Month", "Raw Sales (INR)", "Validated Sales (INR)"], start=1):
            cell = dashboard.cell(trend_header_row, col, header)
            cell.fill = PatternFill("solid", fgColor=TEAL)
            cell.font = Font(name="Aptos", bold=True, color=WHITE)
            cell.alignment = Alignment(wrap_text=True)
        for row_num, trend in enumerate(result["sales_trend"], start=trend_data_row):
            dashboard.cell(row_num, 1, trend["Period"]).number_format = "mmm yyyy" if result["sales_trend_period"] == "Month" else "dd-mmm-yyyy"
            dashboard.cell(row_num, 2, trend["Raw Sales"]).number_format = INR_FORMAT
            dashboard.cell(row_num, 3, trend["Validated Sales"]).number_format = INR_FORMAT
        dashboard["A43"] = f"Sales trend by {result['sales_trend_period'].lower()} (only rows with valid dates)"
        dashboard["A43"].font = Font(name="Aptos", bold=True, color=NAVY, size=11)
        _add_trend_chart(dashboard, trend_data_row, trend_data_row + len(result["sales_trend"]) - 1)

    workbook.active = 0
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def build_profile_excel_report(profile: dict[str, Any], template_result: dict[str, Any] | None = None) -> bytes:
    """Build a generic workbook while retaining source rows and values."""
    validate_profile_schema(profile)
    if template_result is not None:
        validate_template_result_schema(template_result)
    data: pd.DataFrame = profile["original_data"]
    issues: pd.DataFrame = profile["issues"]
    template_result = template_result or {}
    workbook = Workbook()
    dashboard = workbook.active
    dashboard.title = "MIS Dashboard"
    source = workbook.create_sheet("Original Data")
    quality = workbook.create_sheet("Data Quality Issues")
    column_profile = workbook.create_sheet("Column Profile")
    workbook.properties.title = f"MIS Autopilot | {template_result.get('template', 'General')} report"
    workbook.properties.subject = "Spreadsheet profile and data-quality review"
    workbook.properties.creator = "MIS Autopilot"

    title_template = template_result.get("template", "General spreadsheet analysis")
    _write_title(dashboard, "MIS Autopilot | Spreadsheet Report", f"Template: {title_template} · Source: {profile['source_name']} · Original values are preserved.", 10)
    dashboard.merge_cells("A3:D3")
    dashboard["A3"] = "Metrics are calculated only from confirmed column mappings. Generic records and duplicate candidates are retained for review."
    dashboard["A3"].font = Font(name="Aptos", size=10, italic=True, color=TEXT)
    dashboard["A3"].alignment = Alignment(wrap_text=True, vertical="top")
    dashboard.row_dimensions[3].height = 32
    metric_headers = ["Metric", "Value", "Basis / note"]
    for col, value in enumerate(metric_headers, start=1):
        cell = dashboard.cell(5, col, value)
        cell.fill = PatternFill("solid", fgColor=TEAL)
        cell.font = Font(name="Aptos", bold=True, color=WHITE)
    base_metrics = [
        {"label": "Source rows", "value": profile["row_count"], "kind": "number", "note": "Includes blank and suspicious rows."},
        {"label": "Columns", "value": profile["column_count"], "kind": "number", "note": "Duplicate headings are made unique for processing."},
        {"label": "Exact duplicate rows", "value": profile["duplicate_count"], "kind": "number", "note": "Flagged and retained in Original Data."},
        {"label": "Data-quality findings", "value": len(issues), "kind": "number", "note": "A row can have more than one finding."},
    ]
    metrics = base_metrics + list(template_result.get("metrics", []))
    for row_num, item in enumerate(metrics, start=6):
        dashboard.cell(row_num, 1, item["label"])
        dashboard.cell(row_num, 2, _excel_value(item.get("value")))
        dashboard.cell(row_num, 3, item.get("note", ""))
        dashboard.cell(row_num, 1).font = Font(name="Aptos", size=10, color=TEXT)
        dashboard.cell(row_num, 2).font = Font(name="Aptos", size=10, color=TEXT)
        dashboard.cell(row_num, 3).font = Font(name="Aptos", size=10, color=TEXT)
        dashboard.cell(row_num, 3).alignment = Alignment(wrap_text=True, vertical="top")
        if item.get("kind") == "currency":
            dashboard.cell(row_num, 2).number_format = INR_FORMAT
        elif item.get("kind") == "percent":
            dashboard.cell(row_num, 2).number_format = "0.0%"
        elif item.get("kind") == "number":
            dashboard.cell(row_num, 2).number_format = INTEGER_FORMAT if isinstance(item.get("value"), int) else "#,##0.00"
    dashboard.column_dimensions["A"].width = 32
    dashboard.column_dimensions["B"].width = 22
    dashboard.column_dimensions["C"].width = 72
    dashboard.freeze_panes = "A6"

    source_headers = [str(column) for column in data.columns]
    issue_row_nums = set(issues["Row"].astype(int).tolist()) if not issues.empty else set()
    _write_title(source, "Original Data | Source Values", "No source row has been dropped or overwritten. The final columns are review indicators.", max(3, len(source_headers) + 2))
    _write_headers(source, source_headers + ["Quality Status", "Issue Indicators"])
    issue_by_row: dict[int, list[str]] = {}
    for record in issues.to_dict("records"):
        issue_by_row.setdefault(int(record["Row"]), []).append(str(record["Issue"]))
    for index, (_, row) in enumerate(data.iterrows()):
        source_row = index + 2
        row_issues = list(dict.fromkeys(issue_by_row.get(source_row, [])))
        nonblank = any(not _blank(value) for value in row.tolist())
        status = "Needs review" if row_issues else ("Valid" if nonblank else "Blank row")
        values = [_excel_value(value) for value in row.tolist()] + [status, "; ".join(row_issues)]
        excel_row = index + 5
        for col, value in enumerate(values, start=1):
            cell = source.cell(excel_row, col, value)
            cell.font = BODY_FONT
            cell.alignment = BODY_WRAP_ALIGNMENT if col > len(source_headers) else BODY_ALIGNMENT
            cell.border = BODY_BORDER
        source.cell(excel_row, len(source_headers) + 1).fill = STATUS_FILLS[status]
    source_last = max(4 + len(data), 4)
    source_last_col = get_column_letter(max(len(source_headers) + 2, 1))
    source.auto_filter.ref = f"A4:{source_last_col}{source_last}"
    source.column_dimensions[get_column_letter(len(source_headers) + 1)].width = 18
    source.column_dimensions[get_column_letter(len(source_headers) + 2)].width = 42
    if len(source_headers):
        _fit_widths(source, 4, max(source_last, 4), 1, len(source_headers), cap=38)
    source.freeze_panes = "A5"
    source.conditional_formatting.add(
        f"A5:{get_column_letter(max(len(source_headers), 1))}{source_last}",
        FormulaRule(formula=[f'${get_column_letter(len(source_headers) + 1)}5="Needs review"'], fill=PatternFill("solid", fgColor=PALE_AMBER)),
    )

    quality_headers = list(issues.columns)
    _write_title(quality, "Data Quality Issues", "Each finding identifies the source row and keeps the original value visible.", max(1, len(quality_headers)))
    _write_headers(quality, quality_headers)
    for row_num, record in enumerate(issues.itertuples(index=False, name=None), start=5):
        for col, value in enumerate(record, start=1):
            cell = quality.cell(row_num, col, _excel_value(value))
            cell.font = BODY_FONT
            cell.alignment = BODY_WRAP_ALIGNMENT
            cell.border = BODY_BORDER
            if col == quality_headers.index("Severity") + 1:
                severity = str(value).lower()
                cell.fill = PatternFill("solid", fgColor=PALE_RED if severity == "error" else PALE_AMBER)
    quality_last = max(4 + len(issues), 4)
    quality.auto_filter.ref = f"A4:{get_column_letter(max(len(quality_headers), 1))}{quality_last}"
    if quality_headers:
        _fit_widths(quality, 4, max(quality_last, 4), 1, len(quality_headers), cap=60)
    quality.freeze_panes = "A5"

    profile_frame: pd.DataFrame = profile["column_profile"]
    _write_title(column_profile, "Column Profile", "Detected types, completeness, distinct counts and sample values.", max(1, len(profile_frame.columns)))
    _write_headers(column_profile, list(profile_frame.columns))
    for row_num, record in enumerate(profile_frame.itertuples(index=False, name=None), start=5):
        for col, value in enumerate(record, start=1):
            cell = column_profile.cell(row_num, col, _excel_value(value))
            cell.font = BODY_FONT
            cell.alignment = BODY_WRAP_ALIGNMENT
            cell.border = BODY_BORDER
    profile_last = max(4 + len(profile_frame), 4)
    profile_last_col = get_column_letter(max(len(profile_frame.columns), 1))
    column_profile.auto_filter.ref = f"A4:{profile_last_col}{profile_last}"
    if len(profile_frame.columns):
        _fit_widths(column_profile, 4, max(profile_last, 4), 1, len(profile_frame.columns), cap=44)
    column_profile.freeze_panes = "A5"

    summary_sheets: list[tuple[Any, str]] = []
    used_titles = set(workbook.sheetnames)
    for requested_title, frame in template_result.get("summaries", {}).items():
        if not isinstance(frame, pd.DataFrame) or frame.empty:
            continue
        title = str(requested_title)[:31]
        base = title
        counter = 2
        while title in used_titles:
            suffix = f" {counter}"
            title = base[:31 - len(suffix)] + suffix
            counter += 1
        used_titles.add(title)
        ws = workbook.create_sheet(title)
        headers = [str(col) for col in frame.columns]
        _write_title(ws, title, f"Template summary for {title_template}. Calculations use the confirmed mappings.", max(1, len(headers)))
        _write_headers(ws, headers)
        for row_num, values in enumerate(frame.itertuples(index=False, name=None), start=5):
            for col, value in enumerate(values, start=1):
                cell = ws.cell(row_num, col, _excel_value(value))
                cell.font = BODY_FONT
                cell.alignment = BODY_WRAP_ALIGNMENT
                cell.border = BODY_BORDER
        last = max(4 + len(frame), 4)
        ws.auto_filter.ref = f"A4:{get_column_letter(max(len(headers), 1))}{last}"
        if headers:
            _fit_widths(ws, 4, max(last, 4), 1, len(headers), cap=44)
            if len(headers) > 1 and len(frame) > 0:
                chart = BarChart()
                chart.type = "bar"
                chart.style = 10
                chart.title = f"{title} overview"
                chart.y_axis.title = headers[0]
                chart.x_axis.title = headers[1]
                chart.add_data(Reference(ws, min_col=2, max_col=2, min_row=4, max_row=last), titles_from_data=True)
                chart.set_categories(Reference(ws, min_col=1, min_row=5, max_row=last))
                chart.height = 7
                chart.width = 14
                ws.add_chart(chart, f"{get_column_letter(len(headers) + 2)}4")
        ws.freeze_panes = "A5"
        summary_sheets.append((ws, title))

    # Put a compact, template-appropriate chart on the dashboard when a
    # category/numeric summary exists. The detail sheet keeps the full table.
    for ws, title in summary_sheets:
        frame = template_result["summaries"].get(title)
        if frame is None:
            frame = next((value for key, value in template_result["summaries"].items() if str(key)[:31] == title), None)
        if frame is None or frame.empty or len(frame.columns) < 2:
            continue
        preferred = [index for index, header in enumerate(frame.columns, start=1) if str(header).lower() in {"sum", "total", "stock quantity", "records", "actual", "validated sales"}]
        numeric_col = next((index for index in preferred if index > 1), None)
        if numeric_col is None:
            numeric_col = next((index for index in range(2, len(frame.columns) + 1) if pd.api.types.is_numeric_dtype(frame.iloc[:, index - 1])), None)
        if numeric_col is not None:
            chart = BarChart()
            chart.type = "bar"
            chart.style = 10
            chart.title = title
            chart.y_axis.title = str(frame.columns[0])
            chart.x_axis.title = str(frame.columns[numeric_col - 1])
            chart.add_data(Reference(ws, min_col=numeric_col, max_col=numeric_col, min_row=4, max_row=4 + len(frame)), titles_from_data=True)
            chart.set_categories(Reference(ws, min_col=1, min_row=5, max_row=4 + len(frame)))
            chart.height = 8
            chart.width = 16
            dashboard.add_chart(chart, "E5")
            break

    if profile_frame is not None and len(profile_frame) and "Missing values" in profile_frame.columns:
        max_missing = int(pd.to_numeric(profile_frame["Missing values"], errors="coerce").fillna(0).max())
        if max_missing > 0:
            chart = BarChart()
            chart.type = "bar"
            chart.style = 10
            chart.title = "Missing values by column"
            chart.y_axis.title = "Column"
            chart.x_axis.title = "Missing cells"
            chart.add_data(Reference(column_profile, min_col=4, max_col=4, min_row=4, max_row=4 + len(profile_frame)), titles_from_data=True)
            chart.set_categories(Reference(column_profile, min_col=1, min_row=5, max_row=4 + len(profile_frame)))
            chart.height = 8
            chart.width = 15
            dashboard.add_chart(chart, "E22")

    workbook.active = 0
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
