import unittest
from collections import Counter
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from mis_core import analyze_sales, build_report, validate_sales_result_schema
from spreadsheet_core import load_spreadsheet


ROOT = Path(__file__).resolve().parents[1]
REALISTIC_WORKBOOK = ROOT.parent / "mis_autopilot_realistic_sales_10k.xlsx"


class SalesAnalysisTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.sample = pd.read_csv(ROOT / "sample_sales.csv")
        cls.result = analyze_sales(cls.sample, "sample_sales.csv")

    def test_sample_detects_sales_fields_and_duplicate(self):
        self.assertEqual(self.result["mapping"]["order_id"], "Order ID")
        self.assertEqual(self.result["mapping"]["total_sales"], "Total Sales")
        self.assertEqual(self.result["duplicate_count"], 1)
        self.assertEqual(self.result["repeated_transaction_count"], 0)
        self.assertEqual(len(self.result["cleaned_data"]), len(self.sample) - 1)

    def test_sales_result_schema_is_versioned_and_rejects_a_missing_consumer_key(self):
        self.assertEqual(self.result["schema_version"], 1)
        incomplete = dict(self.result)
        del incomplete["repeated_transaction_count"]
        with self.assertRaisesRegex(ValueError, "missing keys: repeated_transaction_count"):
            validate_sales_result_schema(incomplete)

    def test_sample_metrics_exclude_duplicate(self):
        self.assertEqual(self.result["order_count"], 9)
        self.assertEqual(self.result["record_count"], 10)
        self.assertAlmostEqual(self.result["raw_total_sales"], 11960.0)
        self.assertAlmostEqual(self.result["total_sales"], 9800.0)
        self.assertAlmostEqual(self.result["average_order_value"], 9800 / 9)

    def test_quality_flags_expected_problems_with_source_rows(self):
        issues = self.result["issues"]
        kinds = set(issues["Issue"])
        self.assertIn("Invalid date", kinds)
        self.assertIn("Negative quantity", kinds)
        self.assertIn("Negative unit price", kinds)
        self.assertIn("Negative total sales", kinds)
        self.assertIn("Missing product", kinds)
        self.assertIn("Duplicate record", kinds)
        self.assertIn("Total sales does not match quantity × unit price", kinds)
        self.assertTrue((issues["Row"] >= 2).all())

    def test_cleaned_output_and_mis_report(self):
        cleaned = self.result["cleaned_data"]
        self.assertEqual(cleaned.loc[5, "Date"], "")
        report = build_report(self.result)
        self.assertIn("MIS AUTOPILOT | SALES SUMMARY", report)
        self.assertIn("INR 9,800.00", report)
        self.assertIn("TOP PRODUCTS BY VALIDATED SALES", report)

    def test_blank_source_rows_are_preserved_without_shifting_issue_row_numbers(self):
        source = pd.DataFrame([
            {"Order ID": "A1", "Date": "2026-01-01", "Product": "Soap", "Quantity": "2", "Unit Price": "50", "Total Sales": "100"},
            {"Order ID": "", "Date": "", "Product": "", "Quantity": "", "Unit Price": "", "Total Sales": ""},
            {"Order ID": "A2", "Date": "bad-date", "Product": "", "Quantity": "1", "Unit Price": "10", "Total Sales": "10"},
        ])
        result = analyze_sales(source, "with-blank-row.csv")
        self.assertEqual(len(result["original_data"]), 3)
        self.assertEqual(result["record_count"], 2)
        self.assertIn(4, set(result["issues"]["Row"]))

    def test_order_value_ignores_unidentified_sales_rows(self):
        source = pd.DataFrame({
            "Order ID": ["A1", ""],
            "Date": ["2026-01-01", "2026-01-02"],
            "Product": ["Soap", "Oil"],
            "Quantity": [1, 1],
            "Unit Price": [100, 50],
            "Total Sales": [100, 50],
        })
        result = analyze_sales(source)
        self.assertEqual(result["unique_order_count"], 1)
        self.assertEqual(result["validated_total_sales"], 150)
        self.assertEqual(result["aov_sales"], 100)
        self.assertEqual(result["average_order_value"], 100)

    def test_record_id_only_differences_are_candidates_not_exact_duplicates(self):
        source = pd.DataFrame({
            "Record ID": ["REC-1", "REC-2"],
            "Order ID": ["A1", "A1"],
            "Date": ["2026-01-01", "2026-01-01"],
            "Product": ["Soap", "Soap"],
            "Quantity": [2, 2],
            "Unit Price": [50, 50],
            "Total Sales": [100, 100],
            "Region": ["Delhi", "Delhi"],
        })
        result = analyze_sales(source)
        self.assertEqual(result["duplicate_count"], 0)
        self.assertEqual(result["repeated_transaction_count"], 1)
        self.assertEqual(len(result["cleaned_data"]), 2)
        self.assertEqual(result["raw_total_sales"], 200)
        self.assertEqual(result["validated_total_sales"], 200)
        candidate = result["issues"].loc[result["issues"]["Issue"].eq("Repeated transaction candidate")].iloc[0]
        self.assertEqual(candidate["Row"], 3)
        self.assertEqual(candidate["Severity"], "Warning")

    @unittest.skipUnless(REALISTIC_WORKBOOK.is_file(), "realistic uploaded workbook is not beside the project")
    def test_realistic_workbook_duplicate_and_quality_counts(self):
        workbook = load_workbook(REALISTIC_WORKBOOK, read_only=True, data_only=True)
        self.assertEqual(workbook.sheetnames, ["Sales Data", "Dataset Notes"])
        worksheet = workbook["Sales Data"]
        self.assertEqual(worksheet.max_row, 10026)  # header plus 10,025 data rows
        original_rows = list(worksheet.iter_rows(min_row=2, values_only=True))
        exact_counts = Counter(tuple(row) for row in original_rows)
        without_record_id_counts = Counter(tuple(row[1:]) for row in original_rows)
        independently_counted_exact_repeats = sum(count - 1 for count in exact_counts.values() if count > 1)
        independently_counted_transaction_candidates = sum(count - 1 for count in without_record_id_counts.values() if count > 1)
        self.assertEqual(independently_counted_exact_repeats, 0)
        self.assertEqual(independently_counted_transaction_candidates, 25)

        sales = pd.read_excel(REALISTIC_WORKBOOK, sheet_name="Sales Data", dtype=object)
        result = analyze_sales(sales, REALISTIC_WORKBOOK.name)
        self.assertEqual(result["record_count"], 10025)
        self.assertEqual(result["duplicate_count"], 0)
        self.assertEqual(result["repeated_transaction_count"], 25)
        self.assertEqual(len(result["cleaned_data"]), 10025)
        self.assertEqual(result["validated_total_sales"], result["raw_total_sales"])
        self.assertEqual(sum(row["Validated Records"] for row in result["product_summary"]), 10025)
        self.assertTrue(result["original_data"].equals(sales.reset_index(drop=True)))

        # Also exercise the same worksheet-reading path used by the Streamlit app.
        app_loaded, app_sheets = load_spreadsheet(REALISTIC_WORKBOOK.read_bytes(), REALISTIC_WORKBOOK.name, "Sales Data")
        app_result = analyze_sales(app_loaded, REALISTIC_WORKBOOK.name)
        self.assertEqual(app_sheets, ["Sales Data", "Dataset Notes"])
        self.assertEqual(len(app_loaded), 10025)
        self.assertEqual(app_result["duplicate_count"], 0)
        self.assertEqual(app_result["repeated_transaction_count"], 25)
        app_issue_counts = app_result["issues"]["Issue"].value_counts().to_dict()
        for issue_name, expected_count in {
            "Invalid date": 20,
            "Missing product": 28,
            "Missing region": 27,
            "Missing unit price": 27,
            "Negative quantity": 20,
            "Negative unit price": 20,
            "Total sales does not match quantity × unit price": 67,
        }.items():
            self.assertEqual(app_issue_counts[issue_name], expected_count)

        counts = result["issues"]["Issue"].value_counts().to_dict()
        self.assertEqual(counts["Repeated transaction candidate"], 25)
        self.assertEqual(counts["Invalid date"], 20)
        self.assertEqual(counts["Missing product"], 28)
        self.assertEqual(counts["Missing region"], 27)
        self.assertEqual(counts["Missing unit price"], 27)
        self.assertEqual(counts["Negative quantity"], 20)
        self.assertEqual(counts["Negative unit price"], 20)
        self.assertEqual(counts["Total sales does not match quantity \u00d7 unit price"], 67)
        workbook.close()


if __name__ == "__main__":
    unittest.main()
