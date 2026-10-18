import unittest
from io import BytesIO
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from excel_report import build_excel_report
from mis_core import analyze_sales


ROOT = Path(__file__).resolve().parents[1]


class ExcelReportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.sample = pd.read_csv(ROOT / "sample_sales.csv", dtype=object, keep_default_na=False)
        cls.result = analyze_sales(cls.sample, "sample_sales.csv")

    def test_workbook_has_expected_sheets_charts_and_original_values(self):
        content = build_excel_report(self.result)
        workbook = load_workbook(BytesIO(content), data_only=False)
        self.assertEqual(workbook.sheetnames, [
            "MIS Dashboard", "Sales Data", "Data Quality Issues", "Product Summary", "Region Summary",
        ])
        self.assertEqual(len(workbook["MIS Dashboard"]._charts), 3)
        self.assertEqual(workbook["Sales Data"]["B10"].value, "not-a-date")
        self.assertEqual(workbook["Sales Data"]["A13"].value, "ORD-1009")
        self.assertEqual(workbook["Sales Data"]["A14"].value, "ORD-1009")
        self.assertEqual(workbook["Product Summary"]["E9"].value, "=SUM(E5:E8)")
        self.assertEqual(workbook["MIS Dashboard"]["B6"].value, "='Product Summary'!D9")
        self.assertEqual(workbook["MIS Dashboard"]["B7"].value, "='Product Summary'!E9")
        self.assertGreater(len(workbook["Sales Data"].conditional_formatting), 0)
        self.assertGreater(len(workbook["Data Quality Issues"].conditional_formatting), 0)
        issue_headers = [workbook["Data Quality Issues"].cell(4, col).value for col in range(1, 8)]
        self.assertEqual(issue_headers, [
            "Source Row", "Column Name", "Issue Type", "Original Value", "Explanation", "Recommended Action", "Severity",
        ])

    def test_region_sheet_and_trend_chart_are_conditional(self):
        data = pd.DataFrame({
            "Order ID": ["A1", "A2"],
            "Date": ["not-a-date", "also-not-a-date"],
            "Product": ["Soap", "Oil"],
            "Quantity": [2, 1],
            "Unit Price": [50, 100],
            "Total Sales": [100, 100],
        })
        result = analyze_sales(data, "no-region-no-valid-dates.csv")
        workbook = load_workbook(BytesIO(build_excel_report(result)), data_only=False)
        self.assertNotIn("Region Summary", workbook.sheetnames)
        self.assertEqual(len(workbook["MIS Dashboard"]._charts), 1)

    def test_blank_source_rows_stay_in_sales_data(self):
        data = pd.DataFrame([
            {"Order ID": "A1", "Date": "2026-01-01", "Product": "Soap", "Quantity": "1", "Unit Price": "10", "Total Sales": "10"},
            {"Order ID": "", "Date": "", "Product": "", "Quantity": "", "Unit Price": "", "Total Sales": ""},
            {"Order ID": "A2", "Date": "bad-date", "Product": "Oil", "Quantity": "1", "Unit Price": "5", "Total Sales": "5"},
        ])
        result = analyze_sales(data, "blank-line.csv")
        workbook = load_workbook(BytesIO(build_excel_report(result)), data_only=False)
        self.assertEqual(workbook["Sales Data"]["A6"].value, None)
        self.assertEqual(workbook["Sales Data"]["A7"].value, "A2")
        invalid_date_issue = next(
            row for row in workbook["Data Quality Issues"].iter_rows(min_row=5, values_only=True)
            if row[2] == "Invalid date"
        )
        self.assertEqual(invalid_date_issue[0], 4)

    def test_report_marks_id_ignored_candidates_without_dropping_source_rows(self):
        data = pd.DataFrame({
            "Record ID": ["REC-1", "REC-2"],
            "Order ID": ["A1", "A1"],
            "Date": ["2026-01-01", "2026-01-01"],
            "Product": ["Soap", "Soap"],
            "Quantity": [2, 2],
            "Unit Price": [50, 50],
            "Total Sales": [100, 100],
            "Region": ["Delhi", "Delhi"],
        })
        result = analyze_sales(data, "repeated-transaction.csv")
        workbook = load_workbook(BytesIO(build_excel_report(result)), data_only=False)
        sales = workbook["Sales Data"]
        headers = [sales.cell(4, col).value for col in range(1, sales.max_column + 1)]
        repeated_col = headers.index("Repeated Transaction Candidate") + 1
        exact_col = headers.index("Exact Duplicate") + 1
        self.assertEqual(sales.cell(5, exact_col).value, "No")
        self.assertEqual(sales.cell(6, exact_col).value, "No")
        self.assertEqual(sales.cell(6, repeated_col).value, "Yes")
        self.assertEqual(sales.cell(6, 1).value, "REC-2")
        dashboard = workbook["MIS Dashboard"]
        self.assertEqual(dashboard["A11"].value, "Exact duplicate rows")
        self.assertEqual(dashboard["B11"].value, 0)
        self.assertEqual(dashboard["A12"].value, "Repeated transaction candidates")
        self.assertEqual(dashboard["B12"].value, 1)
        self.assertTrue(any(
            row[2] == "Repeated transaction candidate"
            for row in workbook["Data Quality Issues"].iter_rows(min_row=5, values_only=True)
        ))


if __name__ == "__main__":
    unittest.main()
