import unittest
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from excel_report import build_profile_excel_report
from report_templates import analyze_template, default_mappings, mapping_candidates
from spreadsheet_core import profile_spreadsheet


class ReportTemplateTests(unittest.TestCase):
    def test_sales_template_uses_confirmed_fields_and_sales_schema(self):
        data = pd.DataFrame({
            "Invoice ID": ["A1", "A1"], "Item Name": ["Tea", "Tea"],
            "Quantity Sold": [2, 2], "Unit Price": [50, 50], "Line Total": [100, 100],
        })
        analysis = analyze_template(data, "Sales", {
            "order_id": "Invoice ID", "product": "Item Name", "quantity": "Quantity Sold",
            "unit_price": "Unit Price", "total_sales": "Line Total",
        })
        self.assertTrue(analysis["ready"])
        self.assertEqual(analysis["sales_result"]["duplicate_count"], 1)
        self.assertEqual(analysis["sales_result"]["validated_total_sales"], 100)
        self.assertEqual(analysis["sales_result"]["schema_version"], 1)

    def test_sales_and_finance_templates_explain_missing_required_fields(self):
        data = pd.DataFrame({"Employee": ["A"]})
        sales = analyze_template(data, "Sales", {})
        finance = analyze_template(data, "Finance", {})
        self.assertFalse(sales["ready"])
        self.assertIn("Quantity and Unit Price", sales["missing_requirements"][0])
        self.assertIsNone(sales["sales_result"])
        self.assertFalse(finance["ready"])
        self.assertIn("Expense, Actual, or Amount", finance["missing_requirements"][0])

    def test_finance_template_sums_category_and_budget_actual_with_source_rows_included(self):
        data = pd.DataFrame({
            "Expense Amount": ["100", "200", "100"],
            "Category": ["Rent", "Supplies", "Rent"],
            "Budget": ["300", "300", "300"],
            "Actual": ["100", "200", "100"],
        })
        analysis = analyze_template(data, "Finance", {
            "expense": "Expense Amount", "category": "Category", "budget": "Budget", "actual": "Actual",
        })
        self.assertTrue(analysis["ready"])
        metrics = {item["label"]: item["value"] for item in analysis["metrics"]}
        self.assertEqual(metrics["Mapped expense total"], 400)
        self.assertEqual(metrics["Budget total"], 900)
        self.assertEqual(metrics["Actual total"], 400)
        self.assertEqual(analysis["summaries"]["Category Totals"]["Records"].sum(), 3)

    def test_hr_template_reports_headcount_attendance_and_explicit_absences(self):
        data = pd.DataFrame({
            "Employee ID": ["E1", "E2", "E1"],
            "Attendance Status": ["Present", "Absent", "Leave"],
        })
        analysis = analyze_template(data, "HR and attendance", {
            "employee": "Employee ID", "attendance_status": "Attendance Status",
        })
        metrics = {item["label"]: item["value"] for item in analysis["metrics"]}
        self.assertEqual(metrics["Unique employees"], 2)
        self.assertEqual(metrics["Attendance records"], 3)
        self.assertEqual(metrics["Absence records"], 1)

    def test_inventory_template_uses_explicit_threshold_for_low_stock(self):
        data = pd.DataFrame({"SKU": ["X", "Y"], "Stock Quantity": [3, 10], "Reorder Level": [5, 4]})
        analysis = analyze_template(data, "Inventory", {
            "item": "SKU", "stock_quantity": "Stock Quantity", "threshold": "Reorder Level",
        })
        metrics = {item["label"]: item["value"] for item in analysis["metrics"]}
        self.assertEqual(metrics["Total mapped stock units"], 13)
        self.assertEqual(metrics["Low-stock records"], 1)

    def test_operations_template_calculates_actual_against_target_and_optional_summaries(self):
        data = pd.DataFrame({
            "Actual Output": [8, 12], "Target Output": [10, 10],
            "Task Status": ["Done", "Done"], "Owner": ["A", "B"],
        })
        analysis = analyze_template(data, "Operations", {
            "actual": "Actual Output", "target": "Target Output", "status": "Task Status", "owner": "Owner",
        })
        metrics = {item["label"]: item["value"] for item in analysis["metrics"]}
        self.assertEqual(metrics["Actual total"], 20)
        self.assertEqual(metrics["Target total"], 20)
        self.assertEqual(metrics["Target achievement"], 1)
        self.assertIn("Status Summary", analysis["summaries"])
        self.assertIn("Output by Owner", analysis["summaries"])

    def test_general_template_uses_generic_profile_for_object_numeric_and_categories(self):
        data = pd.DataFrame({"Cost": ["20", "30", ""], "Team": ["A", "B", "A"]})
        profile = profile_spreadsheet(data)
        analysis = analyze_template(data, "General spreadsheet analysis", profile=profile)
        self.assertTrue(analysis["ready"])
        self.assertEqual(analysis["metrics"][0]["value"], 3)
        self.assertEqual(analysis["metrics"][2]["value"], 1)
        self.assertIn("Numeric Summary", analysis["summaries"])
        self.assertIn("Team Breakdown", analysis["summaries"])

    def test_ambiguous_mapping_is_not_auto_confirmed(self):
        candidates = mapping_candidates(["Employee ID", "Employee Name"], "HR and attendance")
        self.assertEqual(candidates["employee"], ["Employee ID", "Employee Name"])
        self.assertNotIn("employee", default_mappings(candidates))

    def test_generic_excel_report_has_expected_sheets_chart_and_unchanged_original_values(self):
        data = pd.DataFrame({"Category": ["Rent", "Supplies"], "Amount": [100, 50], "Date": ["2026-01-01", "bad-date"]})
        profile = profile_spreadsheet(data, "finance.csv")
        analysis = analyze_template(data, "General spreadsheet analysis", profile=profile)
        content = build_profile_excel_report(profile, analysis)
        workbook = load_workbook(BytesIO(content), data_only=False)
        self.assertEqual(workbook.sheetnames[:4], ["MIS Dashboard", "Original Data", "Data Quality Issues", "Column Profile"])
        self.assertEqual(workbook["Original Data"]["C6"].value, "bad-date")
        self.assertEqual(workbook["Original Data"].max_row, 6)
        self.assertIn("Invalid date", [row[2].value for row in workbook["Data Quality Issues"].iter_rows(min_row=5)])
        self.assertGreaterEqual(len(workbook["MIS Dashboard"]._charts), 1)


if __name__ == "__main__":
    unittest.main()
